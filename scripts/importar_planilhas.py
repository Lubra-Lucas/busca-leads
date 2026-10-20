"""Normaliza as planilhas de leads da Bia para o formato da tabela public.leads.

Uso:  python3 scripts/importar_planilhas.py <Leads_atualizado.xlsx> <Garnet_Potenciais_Leads.xlsx> <existentes.json> <saida.json>

- Lê só as abas de origem (as abas "Contatos", "TOP 100" e "Resumo" são consolidações sem dado novo).
- Remove duplicatas por @ do Instagram/TikTok, domínio do site e nome normalizado.
- Leads que já existem no banco viram "atualizações" (só preenchem campos vazios).
"""
import json
import os
import re
import sys
import unicodedata

import openpyxl

ND = {None, '', 'N/D', 'n/d', 'Não informado', '-', '—'}


def limpo(v):
    if isinstance(v, str):
        v = v.strip()
    return None if v in ND else v


def sem_acento(s):
    return unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode()


def chave_nome(s):
    s = sem_acento(str(s or '')).lower()
    s = re.sub(r'\(.*?\)', ' ', s)
    s = re.sub(r'\b(oficial|beauty|makeup|make up|cosmeticos|cosmetics|company|brasil|ltda|by|linha|store|loja)\b', ' ', s)
    return re.sub(r'[^a-z0-9]', '', s)


def dominio(url):
    url = limpo(url)
    if not url:
        return None
    d = re.sub(r'^https?://', '', str(url).lower()).split('/')[0]
    d = re.sub(r'^www\.', '', d)
    # links de bio não identificam a empresa
    if d in {'linktr.ee', 'instagram.com', 'tiktok.com', 'wa.me', 'beacons.ai', 'linkin.bio', 'airgo.bio', 'bio.link',
             'facebook.com', 'api.whatsapp.com', 'linkme.bio', 'taplink.cc', 'linkbio.co', 'campsite.bio', 'meulink.com.br',
             'bit.ly', 'trinks.com', 'sites.google.com', 'client.tuaagenda.com', 'app.frizzar.com.br', 'instabio.cc'}:
        return None
    return d


def handle(v, rede='instagram'):
    v = limpo(v)
    if not v:
        return None
    v = str(v).strip()
    m = re.search(rf'{rede}\.com/@?([A-Za-z0-9_.]+)', v)
    if m:
        v = m.group(1)
    v = v.lstrip('@').strip('/').strip()
    if not re.fullmatch(r'[A-Za-z0-9_.]+', v) or v.lower() in HANDLES_DE_PLATAFORMA:
        return None
    return f'@{v.lower()}'


# o scraper do Google Maps às vezes pega o Instagram do site de agendamento, não o do salão
HANDLES_DE_PLATAFORMA = {'bitly', 'trinksoficial', 'p', 'reel', 'explore', 'stories'}


EMAIL = re.compile(r'[\w.+-]+@[\w-]+\.[\w.-]+')
FONE = re.compile(r'(?:\+?55\s*)?\(?\d{2}\)?\s*9?\d{4}[-\s]?\d{4}|0?800[\s-]?\d{3}[\s-]?\d{4}|\d{4}[-\s]\d{4}')


def emails(txt):
    return EMAIL.findall(str(txt or ''))


def eh_celular(fone):
    d = re.sub(r'\D', '', fone or '')
    if d.startswith('55') and len(d) >= 12:
        d = d[2:]
    return len(d) == 11 and d[2] == '9'


def cidade_uf(v):
    v = limpo(v)
    if not v:
        return None, None
    m = re.match(r'(.+?)\s*/\s*([A-Z]{2})\b', str(v))
    if m:
        return m.group(1).strip(), m.group(2)
    return str(v).strip(), None


UF_CIDADE = {'são paulo': 'SP', 'rio de janeiro': 'RJ', 'belo horizonte': 'MG', 'curitiba': 'PR',
             'porto alegre': 'RS', 'brasília': 'DF', 'brasilia': 'DF'}


def linhas(wb, aba, linha_cabecalho):
    rows = list(wb[aba].iter_rows(values_only=True))
    cab = rows[linha_cabecalho - 1]
    for r in rows[linha_cabecalho:]:
        if any(limpo(v) is not None for v in r):
            yield {str(c).strip(): r[i] for i, c in enumerate(cab) if c is not None}


def novo(**k):
    base = dict(tipo='cliente', segmento_id=None, nome=None, subsegmento=None, contato_nome=None, contato_cargo=None,
                cidade=None, uf=None, bairro=None, endereco=None, site=None, instagram=None, tiktok=None,
                seguidores=None, linkedin=None, email=None, telefone=None, whatsapp=None, fonte=None, fonte_url=None,
                score=None, status='novo', observacoes=None, extras={})
    base.update(k)
    base['extras'] = {a: b for a, b in base['extras'].items() if limpo(b) is not None}
    return base


def score_contato(r, base, teto):
    s = base
    if r['contato_nome'] and r['linkedin']:
        s += 15
    if r['email']:
        s += 15
    if r['whatsapp']:
        s += 10
    elif r['telefone']:
        s += 5
    if r['site']:
        s += 5
    return min(s, teto)


# ---------------------------------------------------------------- influenciadoras
def influenciadoras(arq):
    wb = openpyxl.load_workbook(arq, data_only=True)
    out = []
    for r in linhas(wb, 'Leads Beleza Instagram', 1):
        nome = limpo(r['Nome'])
        if not nome:
            continue
        plataforma = limpo(r.get('Plataforma')) or ''
        contato = limpo(r.get('Contato'))
        tipo_contato = (limpo(r.get('Tipo de Contato')) or '').lower()
        link = limpo(r.get('Link do Perfil'))
        h = handle(r.get('@ Influencer'), 'tiktok' if plataforma == 'TikTok' else 'instagram')
        rec = novo(segmento_id='influenciadora', nome=str(nome).strip(),
                   subsegmento=limpo(r.get('Nicho')),
                   seguidores=int(r['Número de Seguidores']) if isinstance(r.get('Número de Seguidores'), (int, float)) else None,
                   fonte='Pesquisa própria (planilha Leads Beleza Instagram)', fonte_url=link,
                   extras={'plataforma': plataforma, 'media_views_reels': r.get('Méd ultimos 12 Reels (desconsiderar reels fixados)'),
                           'preco_reels': r.get('PREÇO REELS'), 'preco_pacote': r.get('PREÇO PACOTE'),
                           'target': (limpo(r.get('Target')) or '').strip() or None, 'tipo_contato': limpo(r.get('Tipo de Contato')),
                           'origem': 'Leads_atualizado.xlsx › Leads Beleza Instagram'})
        if plataforma == 'TikTok':
            rec['tiktok'] = h
        else:
            rec['instagram'] = h
        if contato:
            em = emails(contato)
            if em:
                rec['email'] = em[0].lower()
                if len(em) > 1:
                    rec['extras']['outros_emails'] = ', '.join(em[1:])
            elif 'telefone' in tipo_contato or 'whatsapp' in tipo_contato or FONE.search(str(contato)):
                rec['telefone'] = str(contato)
                if eh_celular(str(contato)) or 'whatsapp' in tipo_contato:
                    rec['whatsapp'] = str(contato)
            elif re.match(r'https?://|www\.|[\w-]+\.(com|br|bio|ee|me)\b', str(contato)):
                rec['extras']['link_bio'] = str(contato)
            else:
                rec['extras']['assessoria'] = str(contato)
            if 'assessoria' in tipo_contato or 'agência' in str(contato).lower():
                rec['extras']['assessoria'] = rec['extras'].get('assessoria') or str(contato)
        if limpo(r.get('envio de mensagem')) == 'sim':
            rec['status'] = 'contatado'
        # score: contato direto pesa mais; Target/Não target vem da avaliação da Bia
        s = 40
        if rec['email']:
            s += 25
        elif rec['whatsapp'] or rec['telefone']:
            s += 20
        elif rec['extras'].get('assessoria'):
            s += 15
        elif rec['extras'].get('link_bio'):
            s += 10
        alvo = (rec['extras'].get('target') or '').lower()
        if alvo == 'target':
            s += 15
        elif alvo == 'não target':
            s -= 20
        seg = rec['seguidores'] or 0
        if 100_000 <= seg < 1_000_000:
            s += 10
        elif seg >= 30_000:
            s += 5
        rec['score'] = max(0, min(100, s))
        if not h and not link:
            rec['subsegmento'] = rec['subsegmento'] or 'Celebridade'
        out.append(rec)

    # respostas de e-mail (aba Monitoramento) viram histórico de contatos
    interacoes = []
    for r in linhas(wb, 'Monitoramento de E-mails', 1):
        nome = limpo(r.get('Nome da Influencer'))
        if not nome:
            continue
        link = limpo(r.get('Link do Perfil\n(TikTok / Instagram)'))
        partes = [limpo(r.get('O que foi dito no E-mail\n(Chamada / Pergunta)'))]
        if limpo(r.get('Nome da Assessora')):
            partes.append(f"Assessoria: {limpo(r['Nome da Assessora'])}")
        if limpo(r.get('Houve Contato?')):
            partes.append(f"Houve contato: {limpo(r['Houve Contato?'])}")
        if limpo(r.get('Vamos Seguir Adiante?')):
            partes.append(f"Seguir adiante: {limpo(r['Vamos Seguir Adiante?'])}")
        interacoes.append(dict(nome=str(nome).strip(), link=link,
                               instagram=handle(link, 'instagram') if link and 'instagram' in link else None,
                               tiktok=handle(link, 'tiktok') if link and 'tiktok' in link else None,
                               telefone=limpo(r.get('Telefone de Contato')),
                               seguir=limpo(r.get('Vamos Seguir Adiante?')),
                               resumo=' · '.join(str(p) for p in partes if p) or 'Respondeu ao e-mail'))
    return out, interacoes


# ---------------------------------------------------------------- empresas
def empresas(arq):
    wb = openpyxl.load_workbook(arq, data_only=True)
    out = []

    def top(aba, segmento):
        for r in linhas(wb, aba, 4):
            cid, uf = cidade_uf(r.get('Cidade / UF'))
            obs = limpo(r.get('Observações'))
            tel = limpo(r.get('Telefone / WhatsApp'))
            rec = novo(segmento_id=segmento, nome=str(r['Empresa / Marca']).strip(), subsegmento=limpo(r.get('Segmento')),
                       cidade=cid, uf=uf, site=limpo(r.get('Site')), instagram=handle(r.get('Instagram')),
                       contato_nome=limpo(r.get('Decisor')), contato_cargo=limpo(r.get('Cargo')),
                       linkedin=limpo(r.get('LinkedIn')), telefone=tel,
                       whatsapp=tel if tel and ('whatsapp' in str(tel).lower() or eh_celular(str(tel))) else None,
                       fonte=limpo(r.get('Fonte do Lead')), score=int(r['Garnet Lead Score']),
                       observacoes=' | '.join(str(x) for x in [limpo(r.get('Motivo do Fit com a Garnet')), obs] if x) or None,
                       extras={'ranking_planilha': r.get('Ranking'), 'porte_funcionarios': limpo(r.get('Porte (funcionários)')),
                               'faturamento_usd': limpo(r.get('Faturamento estimado (USD)')),
                               'origem': f'Garnet_Potenciais_Leads.xlsx › {aba}'})
            em = emails(r.get('E-mail'))
            rec['email'] = em[0].lower() if em else None
            if rec['whatsapp']:
                rec['whatsapp'] = re.sub(r'\s*\(whatsapp\)', '', rec['whatsapp'], flags=re.I)
                rec['telefone'] = rec['whatsapp']
            if obs and re.search(r"pertence ao grupo L'Oréal|sediada em Nova York|não atende ao critério", obs):
                rec['status'] = 'descartado'
            out.append(rec)

    def reserva(aba, segmento, base):
        for r in linhas(wb, aba, 4):
            obs = limpo(r.get('Observação'))
            rec = novo(segmento_id=segmento, nome=str(r['Empresa / Marca']).strip(),
                       subsegmento=limpo(r.get('Segmento (a confirmar)')), site=limpo(r.get('Site')),
                       contato_nome=limpo(r.get('Decisor')), contato_cargo=limpo(r.get('Cargo')),
                       linkedin=limpo(r.get('LinkedIn')), fonte=limpo(r.get('Fonte')), observacoes=obs,
                       extras={'origem': f'Garnet_Potenciais_Leads.xlsx › {aba}'})
            if rec['site'] and not re.match(r'[\w-]+(\.[\w-]+)+', str(rec['site'])):
                rec['site'] = None
            rec['score'] = score_contato(rec, base, 60)
            if obs and obs.startswith('EXCLUÍDO'):
                rec['status'] = 'descartado'
                if 'duplicata' in obs:
                    rec['_duplicata_declarada'] = True
            out.append(rec)

    def salao(r, aba):
        cidade = limpo(r.get('Cidade'))
        end = limpo(r.get('Endereço'))
        m = re.search(r'-\s*([A-Z]{2}),', end or '')
        uf = m.group(1) if m else UF_CIDADE.get((cidade or '').lower())
        tel = limpo(r.get('Telefone'))
        em = emails(r.get('E-mail'))
        rec = novo(segmento_id='salao_beleza', nome=str(r['Salão / Clínica']).strip(), subsegmento=limpo(r.get('Categoria')),
                   cidade=cidade, uf=uf, bairro=limpo(r.get('Bairro')), endereco=end, site=limpo(r.get('Site')),
                   instagram=handle(r.get('Instagram')), telefone=tel, whatsapp=tel if tel and eh_celular(tel) else None,
                   email=em[0].lower() if em else None,
                   contato_nome=limpo(r.get('Decisor (via Clay)')), contato_cargo=limpo(r.get('Cargo')),
                   linkedin=limpo(r.get('LinkedIn')), fonte=limpo(r.get('Fonte do Lead')),
                   observacoes=limpo(r.get('Observações')),
                   extras={'nota_google': r.get('Nota (Google)'), 'avaliacoes_google': r.get('Nº Avaliações'),
                           'ranking_planilha': r.get('Ranking'), 'origem': f'Garnet_Potenciais_Leads.xlsx › {aba}'})
        return rec

    top('Cosméticos', 'marca_indie')
    reserva('Cosméticos - Reserva', 'marca_indie', 30)
    top('Vestuário', 'marca_moda')
    reserva('Vestuário - Reserva', 'marca_moda', 30)
    # salões não são público prioritário da Garnet (revendem, não lançam linha própria): teto 60
    for r in linhas(wb, 'Salões de Beleza', 4):
        rec = salao(r, 'Salões de Beleza')
        rec['score'] = round(float(r['Garnet Lead Score']) * 60)
        out.append(rec)
    for r in linhas(wb, 'Salões de Beleza - Reserva', 4):
        rec = salao(r, 'Salões de Beleza - Reserva')
        rec['score'] = score_contato(rec, 20, 45)
        out.append(rec)
    for r in linhas(wb, 'TikTok Shop e Perfumaria', 4):
        cid, uf = cidade_uf(re.sub(r'\s*\(.*?\)', '', str(r.get('Cidade / UF') or '')))
        rec = novo(segmento_id='tiktok_perfumaria', nome=str(r['Empresa / Marca']).strip(), cidade=cid, uf=uf,
                   site=limpo(r.get('Site')), contato_nome=limpo(r.get('Decisor (via Clay)')),
                   contato_cargo=limpo(r.get('Cargo do Decisor')), linkedin=limpo(r.get('LinkedIn do Decisor')),
                   fonte='Pesquisa própria (planilha Garnet Potenciais Leads)',
                   observacoes=' | '.join(str(x) for x in [limpo(r.get('Motivo do Lead')), limpo(r.get('Observações'))] if x),
                   extras={'porte_funcionarios': limpo(r.get('Porte (funcionários)')),
                           'faturamento_usd': limpo(r.get('Faturamento estimado (USD)')),
                           'origem': 'Garnet_Potenciais_Leads.xlsx › TikTok Shop e Perfumaria'})
        rec['score'] = score_contato(rec, 30, 60)
        out.append(rec)
    return out


# ---------------------------------------------------------------- deduplicação
def chaves(r):
    ks = []
    for campo in ('instagram', 'tiktok'):
        if r.get(campo):
            ks.append(f'{campo}:{r[campo]}')
    d = dominio(r.get('site'))
    if d:
        ks.append(f'site:{d}')
    n = chave_nome(r.get('nome'))
    if n:
        # salões: unidades de rede têm nome parecido; só é o mesmo lead se o telefone bater
        if r.get('segmento_id') == 'salao_beleza':
            fone = re.sub(r'\D', '', r.get('telefone') or '')[-8:]
            ks.append(f"salao:{n}:{fone or chave_nome(r.get('cidade'))}")
        else:
            ks.append(f'nome:{n}')
    return ks


CAMPOS = ['subsegmento', 'contato_nome', 'contato_cargo', 'cidade', 'uf', 'bairro', 'endereco', 'site', 'instagram',
          'tiktok', 'seguidores', 'linkedin', 'email', 'telefone', 'whatsapp', 'fonte_url']


def unidade(r):
    return {k: r.get(k) for k in ('nome', 'endereco', 'bairro', 'cidade', 'telefone', 'site') if r.get(k)}


def mesclar(a, b):
    """b completa a (não sobrescreve)."""
    if a.get('segmento_id') == 'salao_beleza' and b.get('segmento_id') == 'salao_beleza' \
            and chave_nome(b.get('nome')) != chave_nome(a.get('nome')):
        unidades = a.get('extras', {}).get('unidades') or [unidade(a)]
        a.setdefault('extras', {})['unidades'] = unidades + [unidade(b)]
        # nome da rede = parte comum dos nomes das unidades ("Royal Face Clínica de Estética - X")
        comum = os.path.commonprefix([u['nome'] for u in a['extras']['unidades']]).rstrip(' -|–')
        if len(comum) >= 6:
            a['nome'] = f'{comum} (rede)'
    for c in CAMPOS:
        if not a.get(c) and b.get(c):
            a[c] = b[c]
    a['score'] = max(a.get('score') or 0, b.get('score') or 0)
    if b.get('observacoes') and b['observacoes'] not in (a.get('observacoes') or ''):
        a['observacoes'] = ' | '.join(x for x in [a.get('observacoes'), b['observacoes']] if x)
    a['extras'] = {**b.get('extras', {}), **a.get('extras', {})}
    if a.get('status') == 'novo' and b.get('status') not in (None, 'novo'):
        a['status'] = b['status']
    return a


def main():
    arq_inf, arq_emp, arq_exist, saida = sys.argv[1:5]
    infl, interacoes = influenciadoras(arq_inf)
    emp = empresas(arq_emp)
    existentes = json.load(open(arq_exist))

    # apelidos de leads já cadastrados com nome diferente nas planilhas
    APELIDOS = {'bocarosa': 'Boca Rosa Company', 'brunatavares': 'BT Bruna Tavares', 'btbrunatavares': 'BT Bruna Tavares',
                'marimaria': 'Mari Maria Makeup', 'karenbachini': 'Karen Bachini Beauty'}

    indice = {}
    for e in existentes:
        e['_existente'] = True
        for k in chaves(e):
            indice.setdefault(k, e)
    for alias, nome in APELIDOS.items():
        alvo = next((e for e in existentes if e['nome'] == nome), None)
        if alvo:
            indice.setdefault(f'nome:{alias}', alvo)

    novos, atualizados, duplicatas = [], {}, []
    for r in infl + emp:
        if r.pop('_duplicata_declarada', False):
            duplicatas.append((r['nome'], 'marcada como duplicata na própria planilha'))
            continue
        alvo = next((indice[k] for k in chaves(r) if k in indice), None)
        if alvo is None:
            novos.append(r)
            for k in chaves(r):
                indice[k] = r
            continue
        if 'influenciadora' in (alvo.get('segmento_id'), r['segmento_id']) and alvo.get('segmento_id') != r['segmento_id']:
            # a influenciadora e a marca dela são leads distintos, a menos que dividam o mesmo @
            if not any(k.startswith(('instagram:', 'tiktok:')) for k in set(chaves(r)) & set(chaves(alvo))):
                novos.append(r)
                continue
        duplicatas.append((r['nome'], f"= {alvo['nome']} ({'banco' if alvo.get('_existente') else r['extras'].get('origem')})"))
        if alvo.get('_existente'):
            atual = atualizados.setdefault(alvo['id'], {'id': alvo['id'], 'nome': alvo['nome'], 'extras': {}})
            mesclar(atual, r)
        else:
            mesclar(alvo, r)

    for r in novos:
        r.pop('_existente', None)
    json.dump({'novos': novos, 'atualizados': list(atualizados.values()), 'interacoes': interacoes,
               'duplicatas': duplicatas}, open(saida, 'w'), ensure_ascii=False, default=str, indent=1)
    print(f'influenciadoras lidas: {len(infl)} | empresas lidas: {len(emp)}')
    print(f'novos: {len(novos)} | atualizam existentes: {len(atualizados)} | duplicatas removidas: {len(duplicatas)}')


if __name__ == '__main__':
    main()
